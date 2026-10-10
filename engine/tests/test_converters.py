"""Meaningful conversions preserve source fidelity and produce managed editable versions."""
import hashlib
import io
import json
from pathlib import Path
import re
import shutil
import tempfile
import unittest

from lumen.converters import convert, extract_pptx_text
from lumen.store import Store
from lumen.types import LumenError


def ppt_fixture(path):
    canonical = Path(__file__).resolve().parents[2] / "Resources/Acceptance/reader.pptx"
    if canonical.is_file():
        import shutil
        shutil.copy2(canonical, path)
        return
    # This is a fixture authoring dependency; the production converter uses stdlib + Pillow.
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[6])
    box = first.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(6), Inches(1))
    box.text_frame.text = "真实页序第一章"
    box.text_frame.add_paragraph().text = "原始说明"
    table = first.shapes.add_table(2, 2, Inches(0.5), Inches(1.6), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "能力"; table.cell(0, 1).text = "结果"
    table.cell(1, 0).text = "编辑"; table.cell(1, 1).text = "可保存"
    image = Image.new("RGB", (32, 24), "#5079a0")
    data = io.BytesIO(); image.save(data, format="PNG"); data.seek(0)
    first.shapes.add_picture(data, Inches(6), Inches(0.5), width=Inches(1))
    deck.slides.add_slide(deck.slide_layouts[6])  # Blank page must survive.
    third = deck.slides.add_slide(deck.slide_layouts[6])
    third.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(6), Inches(1)).text_frame.text = "第三章末尾"
    deck.save(path)


class ConversionTests(unittest.TestCase):
    def test_ppt_preserves_page_order_blank_page_table_image_and_durable_edit(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary).resolve(); source = root / "课件.pptx"; ppt_fixture(source)
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            text = extract_pptx_text(source)
            self.assertIn("[第 1 页]", text)
            self.assertIn("[第 3 页]", text)
            self.assertNotIn("[第 2 页]", text)
            self.assertLess(text.index("真实页序第一章"), text.index("第三章末尾"))
            self.assertIn("可保存", text)
            result = convert(root / "state", "ppt", source)
            self.assertTrue(result["validation"]["ok"], result["validation"])
            self.assertEqual(result["conversion"]["pages"], 3)
            self.assertEqual(result["conversion"]["images"], 1)
            self.assertEqual(result["document"]["source_path"], str(source))
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)
            page = Path(result["document"]["entry_path"]).read_text()
            self.assertEqual(len(re.findall(r'<section class="slide', page)), 3)
            self.assertIn("data:image/webp;base64,", page)
            self.assertIn("可保存", page)
            self.assertIn('contenteditable="true"', page)
            init = json.loads(re.search(r'id="html-feedback-init" type="application/json">(.*?)</script>', page, re.S)[1])
            self.assertTrue(init["native_editor"])
            self.assertEqual(init["document_id"], result["document"]["id"])
            modified = re.sub(r'(contenteditable="true">)原始说明(</(?:p|div)>)', r'\1人工补充的说明\2', page, count=1)
            self.assertNotEqual(modified, page)
            store = Store(root / "state")
            document = result["document"]; current = store.get(document["id"])
            feedback = {"page_id": current["pages"][0]["id"], "comments": [{"id":"u1","block_id":"slide-1-block-2","body":"保留这条建议"}], "blocks": {}}
            saved = store.save_document(document["id"], current["pages"][0]["id"], 1, [], feedback, modified)
            reopened = Path(saved["document"]["entry_path"]).read_text()
            self.assertRegex(reopened, r'>人工补充的说明</(?:p|div)>')
            self.assertIn("data:image/webp;base64,", reopened)
            self.assertEqual(store.get(document["id"])["feedback"][current["pages"][0]["id"]]["comments"][0]["body"], "保留这条建议")
            native = json.loads(re.search(r'id="html-feedback-init" type="application/json">(.*?)</script>', reopened, re.S)[1])
            self.assertTrue(native["native_editor"])
            self.assertEqual(native["version"], 2)

    @unittest.skipUnless(shutil.which("folio"), "optional Folio integration")
    def test_real_folio_directory_graph_uses_existing_engine_and_imports_complete_html(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary).resolve(); source = root / "资料"; source.mkdir()
            (source / "项目.md").write_text("# 项目\n\n[依据](依据.md)\n\n解释当前成果。")
            (source / "依据.md").write_text("# 依据\n\n这是关系的来源。")
            before = {path.name:path.read_bytes() for path in source.iterdir()}
            result = convert(root / "state", "graph", source)
            self.assertTrue(result["validation"]["ok"], result["validation"])
            self.assertEqual(result["conversion"]["engine"], "Folio")
            self.assertEqual(result["conversion"]["report"]["files"], 2)
            self.assertGreaterEqual(result["conversion"]["report"]["nodes"], 3)
            self.assertEqual({path.name:path.read_bytes() for path in source.iterdir()}, before)
            page = Path(result["document"]["entry_path"]).read_text()
            self.assertIn('id="lumen-graph-adapter"', page)
            self.assertIn("lumen_overrides", page)
            self.assertIn("window.LumenReader", page)
            self.assertIn("项目.md", page)
            self.assertTrue(Path(result["conversion"]["audit_path"]).is_file())
            store = Store(root / "state")
            self.assertEqual(len(store.list()["documents"]), 1)
            self.assertEqual(store.document(result["document"]["id"])["source_path"], str(source))

    def test_failed_conversion_does_not_register_a_finished_document(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary).resolve()
            with self.assertRaises(LumenError):
                convert(root / "state", "ppt", root / "missing.pptx")
            invalid = root / "broken.pptx"; invalid.write_bytes(b"not a zip")
            with self.assertRaises(LumenError):
                convert(root / "state", "ppt", invalid)
            self.assertEqual(Store(root / "state").list()["documents"], [])


if __name__ == "__main__":
    unittest.main()
